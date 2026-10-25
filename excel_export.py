"""
DealScope India — Excel Report Generator Module
Creates a professionally formatted multi-sheet Excel workbook with 10 dedicated sheets:
1. Executive Dashboard
2. Raw Deals
3. Bulk Deals
4. Block Deals
5. Company Summary
6. Client Summary
7. Daily Summary
8. Exchange Summary
9. Price Reactions
10. Data Quality
"""
from datetime import datetime
import io
from typing import Any, Dict, List, Optional

from analytics import (
    aggregate_by_date,
    aggregate_by_deal_type,
    aggregate_by_exchange,
    compute_market_kpis,
    get_top_clients,
    get_top_companies,
)

def create_excel_report(deals: List[Dict[str, Any]], price_reactions: Optional[List[Dict[str, Any]]] = None) -> io.BytesIO:
    """
    Builds a 10-sheet Excel workbook as an in-memory BytesIO buffer.
    Uses openpyxl with styles, frozen headers, and auto-fit column widths.
    """
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    # remove default sheet
    wb.remove(wb.active)

    # Styles
    navy_header_fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    accent_header_fill = PatternFill(start_color="2B6CB0", end_color="2B6CB0", fill_type="solid")
    sub_header_fill = PatternFill(start_color="EDF2F7", end_color="EDF2F7", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    bold_font = Font(name="Calibri", size=11, bold=True)
    regular_font = Font(name="Calibri", size=10)
    title_font = Font(name="Calibri", size=16, bold=True, color="1B365D")
    kpi_val_font = Font(name="Calibri", size=14, bold=True, color="0D47A1")

    thin_border = Border(
        left=Side(style="thin", color="D1D5DB"),
        right=Side(style="thin", color="D1D5DB"),
        top=Side(style="thin", color="D1D5DB"),
        bottom=Side(style="thin", color="D1D5DB")
    )

    kpis = compute_market_kpis(deals)
    bulk_deals = [d for d in deals if d.get("deal_type") == "BULK"]
    block_deals = [d for d in deals if d.get("deal_type") == "BLOCK"]
    top_comps = get_top_companies(deals, top_n=50)
    top_clis = get_top_clients(deals, top_n=50)
    daily_summary = aggregate_by_date(deals)
    exchange_summary = aggregate_by_exchange(deals)

    def style_headers(ws, row_idx=1, fill=navy_header_fill):
        ws.row_dimensions[row_idx].height = 24
        for cell in ws[row_idx]:
            cell.fill = fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def auto_fit_columns(ws):
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or "")
                if len(val) > max_len:
                    max_len = len(val)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # ---------------- 1. Executive Dashboard ----------------
    ws1 = wb.create_sheet(title="Executive Dashboard")
    ws1.views.sheetView[0].showGridLines = True
    ws1.merge_cells("A1:G1")
    ws1["A1"] = "DealScope India — Executive Bulk & Block Deals Intelligence Dashboard"
    ws1["A1"].font = title_font
    ws1["A2"] = f"Report Generated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M:%S UTC')} | Total Disclosed Trades: {len(deals)}"
    ws1["A2"].font = Font(name="Calibri", size=10, italic=True)

    # KPI Block
    kpi_cards = [
        ("Total Disclosed Deals", kpis["total_deals"], "#,##0"),
        ("Unique Companies", kpis["unique_companies"], "#,##0"),
        ("Unique Disclosed Clients", kpis["unique_clients"], "#,##0"),
        ("Total Disclosed Value (INR)", kpis["total_effective_value"], "₹#,##0.00"),
        ("Reported Deal Value (INR)", kpis["total_reported_value"], "₹#,##0.00"),
        ("Calculated Deal Value (INR)", kpis["total_calculated_value"], "₹#,##0.00"),
        ("Buy-Side Trades", kpis["buy_deals_count"], "#,##0"),
        ("Sell-Side Trades", kpis["sell_deals_count"], "#,##0"),
        ("Total Traded Volume (Shares)", kpis["buy_volume"] + kpis["sell_volume"], "#,##0"),
    ]
    row_curr = 4
    for title, val, num_fmt in kpi_cards:
        ws1.cell(row=row_curr, column=1, value=title).font = bold_font
        ws1.cell(row=row_curr, column=1).fill = sub_header_fill
        c_val = ws1.cell(row=row_curr, column=2, value=val)
        c_val.font = kpi_val_font
        c_val.number_format = num_fmt
        row_curr += 1

    ws1.column_dimensions["A"].width = 30
    ws1.column_dimensions["B"].width = 25

    # ---------------- 2. Raw Deals ----------------
    ws2 = wb.create_sheet(title="Raw Deals")
    deal_headers = [
        "Deal ID", "Trade Date", "Exchange", "Deal Type", "Symbol", "Company Name",
        "ISIN", "Client Name", "Buy / Sell", "Quantity", "Price (INR)",
        "Deal Value (Reported)", "Deal Value (Calculated)", "Deal Value (Effective)",
        "Is Calculated?", "Currency", "Data Source", "Source URL", "Retrieval Timestamp", "Validation Flags"
    ]
    ws2.append(deal_headers)
    for d in deals:
        ws2.append([
            d.get("id"), d.get("trade_date"), d.get("exchange"), d.get("deal_type"), d.get("symbol"),
            d.get("company_name"), d.get("isin") or "N/A", d.get("client_name"), d.get("buy_sell"),
            d.get("quantity"), d.get("price"), d.get("deal_value_reported") or "", d.get("deal_value_calculated"),
            d.get("deal_value_effective"), "YES" if d.get("is_value_calculated") else "NO",
            d.get("currency", "INR"), d.get("data_source_type"), d.get("source_url"),
            d.get("retrieval_timestamp"), d.get("validation_flags")
        ])
    style_headers(ws2)
    auto_fit_columns(ws2)

    # ---------------- 3. Bulk Deals ----------------
    ws3 = wb.create_sheet(title="Bulk Deals")
    ws3.append(deal_headers)
    for d in bulk_deals:
        ws3.append([
            d.get("id"), d.get("trade_date"), d.get("exchange"), d.get("deal_type"), d.get("symbol"),
            d.get("company_name"), d.get("isin") or "N/A", d.get("client_name"), d.get("buy_sell"),
            d.get("quantity"), d.get("price"), d.get("deal_value_reported") or "", d.get("deal_value_calculated"),
            d.get("deal_value_effective"), "YES" if d.get("is_value_calculated") else "NO",
            d.get("currency", "INR"), d.get("data_source_type"), d.get("source_url"),
            d.get("retrieval_timestamp"), d.get("validation_flags")
        ])
    style_headers(ws3)
    auto_fit_columns(ws3)

    # ---------------- 4. Block Deals ----------------
    ws4 = wb.create_sheet(title="Block Deals")
    ws4.append(deal_headers)
    for d in block_deals:
        ws4.append([
            d.get("id"), d.get("trade_date"), d.get("exchange"), d.get("deal_type"), d.get("symbol"),
            d.get("company_name"), d.get("isin") or "N/A", d.get("client_name"), d.get("buy_sell"),
            d.get("quantity"), d.get("price"), d.get("deal_value_reported") or "", d.get("deal_value_calculated"),
            d.get("deal_value_effective"), "YES" if d.get("is_value_calculated") else "NO",
            d.get("currency", "INR"), d.get("data_source_type"), d.get("source_url"),
            d.get("retrieval_timestamp"), d.get("validation_flags")
        ])
    style_headers(ws4)
    auto_fit_columns(ws4)

    # ---------------- 5. Company Summary ----------------
    ws5 = wb.create_sheet(title="Company Summary")
    ws5.append(["Symbol", "Company Name", "Disclosed Deals Count", "Total Deal Value (INR)", "Buy Value (INR)", "Sell Value (INR)", "Total Shares Volume"])
    for c in top_comps:
        ws5.append([c["symbol"], c["company_name"], c["deal_count"], c["total_value"], c["buy_value"], c["sell_value"], c["volume"]])
    style_headers(ws5)
    auto_fit_columns(ws5)

    # ---------------- 6. Client Summary ----------------
    ws6 = wb.create_sheet(title="Client Summary")
    ws6.append(["Client / Investor Name", "Total Deals", "Buy Trades", "Sell Trades", "Unique Companies", "Total Value (INR)", "Buy Value (INR)", "Sell Value (INR)", "Volume"])
    for cl in top_clis:
        ws6.append([cl["client_name"], cl["deal_count"], cl["buy_count"], cl["sell_count"], cl["unique_companies_count"], cl["total_value"], cl["buy_value"], cl["sell_value"], cl["volume"]])
    style_headers(ws6)
    auto_fit_columns(ws6)

    # ---------------- 7. Daily Summary ----------------
    ws7 = wb.create_sheet(title="Daily Summary")
    ws7.append(["Trade Date", "Deal Count", "Total Value (INR)", "Buy Value (INR)", "Sell Value (INR)", "Shares Volume"])
    for dy in daily_summary:
        ws7.append([dy["date"], dy["count"], dy["total_value"], dy["buy_value"], dy["sell_value"], dy["volume"]])
    style_headers(ws7)
    auto_fit_columns(ws7)

    # ---------------- 8. Exchange Summary ----------------
    ws8 = wb.create_sheet(title="Exchange Summary")
    ws8.append(["Exchange", "Trade Count", "Total Disclosed Value (INR)", "Total Shares Volume"])
    for ex in exchange_summary:
        ws8.append([ex["exchange"], ex["count"], ex["total_value"], ex["volume"]])
    style_headers(ws8)
    auto_fit_columns(ws8)

    # ---------------- 9. Price Reactions ----------------
    ws9 = wb.create_sheet(title="Price Reactions")
    ws9.append(["Symbol", "Deal Date", "Session Date", "Deal Price (INR)", "T0 Close (INR)", "Premium/Discount %", "T+1 Close", "T+1 Return %", "T+5 Close", "T+5 Return %", "T+20 Close", "T+20 Return %", "Methodology / Warning"])
    if price_reactions:
        for pr in price_reactions:
            ws9.append([
                pr.get("symbol"), pr.get("deal_date"), pr.get("matched_session_date"), pr.get("deal_price"),
                pr.get("t0_close"), pr.get("premium_to_close_pct"), pr.get("t1_close"), pr.get("t1_return_pct"),
                pr.get("t5_close"), pr.get("t5_return_pct"), pr.get("t20_close"), pr.get("t20_return_pct"),
                "Historical observation only; not indicative of causality or trading performance."
            ])
    else:
        ws9.append(["(No price reaction calculations selected in current export filter)", "", "", "", "", "", "", "", "", "", "", "", "Note: Configure symbol and dates in Price Reactions page."])
    style_headers(ws9)
    auto_fit_columns(ws9)

    # ---------------- 10. Data Quality ----------------
    ws10 = wb.create_sheet(title="Data Quality")
    ws10.append(["Metric / Check", "Status / Count", "Audit Details"])
    ws10.append(["Total Deals Ingested", len(deals), "Verified records"])
    ws10.append(["Calculated Values Count", kpis["calculated_value_deals_count"], "Value derived from Quantity * Price"])
    ws10.append(["Reported Values Count", kpis["reported_value_deals_count"], "Directly disclosed by exchange"])
    ws10.append(["Zero / Negative Price Flags", sum(1 for d in deals if "Suspicious" in d.get("validation_flags", "")), "Anomaly check"])
    ws10.append(["Missing ISIN Records", sum(1 for d in deals if not d.get("isin")), "Exchange omitted ISIN field"])
    style_headers(ws10)
    auto_fit_columns(ws10)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
